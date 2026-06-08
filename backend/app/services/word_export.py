from __future__ import annotations

from datetime import date
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def _format_money(value: float) -> str:
    return f"¥{value:.2f}"


def build_monthly_work_report_docx(
    *,
    employee_name: str,
    employee_account: str,
    month: str,
    export_date: date,
    rows: list[dict[str, object]],
    summary: dict[str, object],
) -> bytes:
    document = Document()

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("员工月度报工详情")
    title_run.bold = True
    title_run.font.size = Pt(18)

    document.add_paragraph(f"员工姓名：{employee_name}")
    document.add_paragraph(f"员工账号：{employee_account}")
    document.add_paragraph(f"导出月份：{month}")
    document.add_paragraph(f"导出时间：{export_date.isoformat()}")
    document.add_paragraph("")

    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headers = ("日期", "工序", "数量", "单价", "工资", "状态")
    for index, label in enumerate(headers):
        table.rows[0].cells[index].text = label

    if rows:
        previous_date = ""
        daily_wage_map: dict[str, float] = {}
        for row in rows:
            work_date = str(row.get("workDate") or "")
            date_total = float(row.get("dateTotalPrice") or 0)
            if date_total > 0:
                daily_wage_map[work_date] = date_total
            elif work_date not in daily_wage_map:
                daily_wage_map[work_date] = float(row.get("totalPrice") or 0)
            else:
                daily_wage_map[work_date] += float(row.get("totalPrice") or 0)

        for index, row in enumerate(rows):
            work_date = str(row.get("workDate") or "")
            show_date = work_date if work_date != previous_date else ""
            cells = table.add_row().cells
            cells[0].text = show_date
            cells[1].text = str(row.get("processName") or "")
            cells[2].text = f"{float(row.get('quantity') or 0):.2f}"
            cells[3].text = _format_money(float(row.get("unitPrice") or 0))
            cells[4].text = _format_money(float(row.get("totalPrice") or 0))
            cells[5].text = str(row.get("confirmStatus") or "未确认")
            previous_date = work_date

            next_date = str(rows[index + 1].get("workDate") or "") if index + 1 < len(rows) else ""
            if work_date and work_date != next_date:
                total_row = table.add_row().cells
                total_row[0].text = ""
                total_row[1].text = "当天总工资"
                total_row[2].text = ""
                total_row[3].text = ""
                total_row[4].text = _format_money(daily_wage_map.get(work_date, 0))
                total_row[5].text = ""
    else:
        empty_row = table.add_row().cells
        empty_row[0].text = "暂无报工记录"
        for index in range(1, 6):
            empty_row[index].text = ""

    document.add_paragraph("")
    document.add_paragraph(f"月工资总额：{_format_money(float(summary.get('totalWage') or 0))}")
    document.add_paragraph(f"上班天数：{int(summary.get('workDays') or 0)}")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
